"""Office Document Agent (L2, "Document" specialist alongside PDF) --
DOCX/PPTX via python-docx/python-pptx (pure-Python, no heavy deps), plus
legacy binary DOC/PPT (pre-2007 OLE2 format) and RTF/ODT via a headless
LibreOffice conversion to the modern XML format first. There is no
reliable pure-Python reader for any of those older formats, and a
text-only extractor (antiword/catdoc) would silently drop every table --
converting first lets all of them reuse the exact same
_parse_docx/_parse_pptx logic below, tables and all.
"""
import asyncio
import logging
import shutil
import subprocess
import tempfile
from pathlib import Path

import docx
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.models.schemas import FileCategory, ParsedDocument, TableBlock, TextBlock
from app.parsers.base import BaseParser

logger = logging.getLogger(__name__)

LIBREOFFICE_TIMEOUT_S = 120

# Office Open XML files (.docx/.pptx) are ZIP archives -- this is the ZIP
# local-file-header magic. A password-protected .docx/.pptx is instead
# stored as an OLE2 compound file (the same container format as the old
# binary .doc/.ppt), which is what Word/PowerPoint use to wrap the whole
# encrypted package. Checking for this up front turns a confusing
# "Package not found" exception into a clear, actionable warning.
_ZIP_MAGIC = b"PK\x03\x04"


class OfficeParser(BaseParser):
    category = FileCategory.OFFICE

    async def parse(self, file_path: str) -> ParsedDocument:
        return await asyncio.to_thread(self._parse_sync, file_path)

    def _parse_sync(self, file_path: str) -> ParsedDocument:
        doc = ParsedDocument(source_file=file_path, category=self.category)
        lower = file_path.lower()
        tmp_dir: Path | None = None
        try:
            if lower.endswith((".docx", ".pptx")) and self._looks_encrypted(file_path):
                doc.warnings.append(
                    "This file appears to be password-protected (it's an OLE2 container, not a "
                    "ZIP archive, which is how Word/PowerPoint wrap an encrypted .docx/.pptx). "
                    "Password-protected Office files aren't supported -- remove the password and "
                    "re-upload."
                )
                return doc
            if lower.endswith(".ppt"):
                converted, tmp_dir = self._convert_legacy(file_path, "pptx")
                self._parse_pptx(converted, doc)
            elif lower.endswith((".doc", ".rtf", ".odt")):
                converted, tmp_dir = self._convert_legacy(file_path, "docx")
                self._parse_docx(converted, doc)
            elif lower.endswith(".pptx"):
                self._parse_pptx(file_path, doc)
            else:
                self._parse_docx(file_path, doc)
        except Exception as exc:  # noqa: BLE001
            logger.exception("Office parse failed on %s", file_path)
            doc.warnings.append(f"Office document parse error: {exc}")
        finally:
            if tmp_dir is not None:
                shutil.rmtree(tmp_dir, ignore_errors=True)
        return doc

    @staticmethod
    def _looks_encrypted(file_path: str) -> bool:
        try:
            with open(file_path, "rb") as f:
                head = f.read(4)
        except OSError:
            return False
        return head != _ZIP_MAGIC

    def _convert_legacy(self, file_path: str, target_ext: str) -> tuple[str, Path]:
        """Converts a legacy .doc/.ppt/.rtf/.odt file to .docx/.pptx via
        headless LibreOffice. Returns the converted file's path and the
        temp output directory it lives in (caller's responsibility to
        clean up).

        A unique -env:UserInstallation profile dir is required per call --
        soffice's default profile is a single shared, lock-file-guarded
        directory, so two conversions running concurrently (this pipeline
        processes up to MAX_PARALLEL_FILES files at once) would otherwise
        serialize on that lock or fail outright with "soffice already
        running; exiting."."""
        outdir = Path(tempfile.mkdtemp(prefix="office_convert_out_"))
        profile_dir = Path(tempfile.mkdtemp(prefix="office_convert_profile_"))
        try:
            result = subprocess.run(
                [
                    "soffice", "--headless", "--norestore",
                    f"-env:UserInstallation=file://{profile_dir}",
                    "--convert-to", target_ext, "--outdir", str(outdir), file_path,
                ],
                capture_output=True, text=True, timeout=LIBREOFFICE_TIMEOUT_S,
            )
            converted = outdir / f"{Path(file_path).stem}.{target_ext}"
            if result.returncode != 0 or not converted.exists():
                raise RuntimeError(
                    "LibreOffice conversion to ." + target_ext + " failed: "
                    + (result.stderr.strip() or result.stdout.strip() or f"exit code {result.returncode}")
                )
            return str(converted), outdir
        except Exception:
            shutil.rmtree(outdir, ignore_errors=True)
            raise
        finally:
            shutil.rmtree(profile_dir, ignore_errors=True)

    def _parse_docx(self, file_path: str, doc: ParsedDocument) -> None:
        d = docx.Document(file_path)
        for para in d.paragraphs:
            if not para.text.strip():
                continue
            kind = "heading" if para.style and para.style.name and para.style.name.startswith("Heading") else "paragraph"
            doc.text_blocks.append(TextBlock(text=para.text, kind=kind))

        for table in d.tables:
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            headers = rows[0] if rows else []
            doc.tables.append(TableBlock(headers=headers, rows=rows[1:]))

        # Headers/footers often carry real signal (classification banners,
        # doc IDs, confidentiality notices, company letterhead) that
        # d.paragraphs never sees -- it only walks the main document body.
        # is_linked_to_previous skips a section whose header/footer is
        # just inherited from the one before it, so a multi-section
        # document with one consistent header doesn't get it duplicated
        # once per section.
        header_footer_count = 0
        for section in d.sections:
            if not section.header.is_linked_to_previous:
                for para in section.header.paragraphs:
                    if para.text.strip():
                        doc.text_blocks.append(TextBlock(text=para.text, kind="header"))
                        header_footer_count += 1
            if not section.footer.is_linked_to_previous:
                for para in section.footer.paragraphs:
                    if para.text.strip():
                        doc.text_blocks.append(TextBlock(text=para.text, kind="footer"))
                        header_footer_count += 1

        doc.metadata = {
            "paragraph_count": len(d.paragraphs),
            "table_count": len(d.tables),
            "header_footer_block_count": header_footer_count,
            "parser": "python-docx",
        }
        if not doc.text_blocks and not doc.tables:
            doc.warnings.append("No extractable text or tables found in document.")

    def _parse_pptx(self, file_path: str, doc: ParsedDocument) -> None:
        prs = Presentation(file_path)
        chart_count = 0
        for i, slide in enumerate(prs.slides, start=1):
            for shape in self._iter_shapes(slide.shapes):
                if shape.has_text_frame and shape.text_frame.text.strip():
                    doc.text_blocks.append(
                        TextBlock(text=shape.text_frame.text, page=i, kind="paragraph")
                    )
                if shape.has_table:
                    rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                    headers = rows[0] if rows else []
                    doc.tables.append(TableBlock(page=i, headers=headers, rows=rows[1:]))
                if shape.has_chart:
                    chart_table = self._chart_to_table(shape, i)
                    if chart_table is not None:
                        doc.tables.append(chart_table)
                        chart_count += 1
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                doc.text_blocks.append(
                    TextBlock(text=slide.notes_slide.notes_text_frame.text, page=i, kind="notes")
                )

        doc.metadata = {"slide_count": len(prs.slides), "chart_count": chart_count, "parser": "python-pptx"}
        if not doc.text_blocks and not doc.tables:
            doc.warnings.append("No extractable text or tables found in presentation.")

    @classmethod
    def _iter_shapes(cls, shapes):
        """Recurses into grouped shapes -- python-pptx's slide.shapes only
        lists top-level shapes, so text/tables/charts a user grouped
        together in PowerPoint (a common way to build diagrams, labeled
        org charts, etc.) would otherwise be invisible to this parser
        entirely, since a GroupShape itself has no text_frame/table/chart
        of its own -- only its children do."""
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from cls._iter_shapes(shape.shapes)
            else:
                yield shape

    @staticmethod
    def _chart_to_table(shape, page: int) -> TableBlock | None:
        """Best-effort: python-pptx's chart data API differs somewhat
        across chart types (pie vs bar vs line vs scatter), and a chart
        with no plots/series is a real, unremarkable case (an empty
        placeholder). Any failure here is one shape's chart, not the
        whole file, so it degrades to skipping just that chart rather
        than failing the presentation."""
        try:
            chart = shape.chart
            plot = chart.plots[0]
            categories = [str(c) for c in plot.categories]
            rows = []
            for series in plot.series:
                values = [("" if v is None else str(v)) for v in series.values]
                rows.append([series.name or ""] + values)
            if not rows:
                return None
            caption = None
            if chart.has_title and chart.chart_title.has_text_frame:
                caption = chart.chart_title.text_frame.text or None
            return TableBlock(page=page, headers=["Series", *categories], rows=rows, caption=caption)
        except Exception:  # noqa: BLE001
            logger.debug("Skipped unreadable chart on slide %s", page, exc_info=True)
            return None
